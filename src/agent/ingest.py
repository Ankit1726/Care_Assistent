import argparse
import hashlib
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from src.agent.knowledge import embed, ensure_index  # noqa: E402

SUPPORTED = {".pdf", ".txt", ".md", ".docx", ".csv"}
PAGE_CHARS = 3000  # pseudo-"page" size for formats without real pages
BATCH = 50


def chunks(text, size=1100, overlap=180):
    if size <= 0 or overlap < 0 or overlap >= size:
        raise ValueError("Chunk size must be positive and overlap smaller than size")
    text = " ".join(text.split())
    step = size - overlap
    return [
        text[i : i + size]
        for i in range(0, len(text), step)
        if len(text[i : i + size]) > 80
    ]


def read_pages(path: Path):
    """Yield (page_number, text) for any supported file."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        import fitz  # PyMuPDF

        with fitz.open(path) as doc:
            for n, page in enumerate(doc, start=1):
                yield n, page.get_text()
        return
    if ext in (".txt", ".md", ".csv"):
        text = path.read_text(encoding="utf-8", errors="ignore")
    elif ext == ".docx":
        from docx import Document  # python-docx

        d = Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError(
            f"Unsupported file type '{ext}'. Use: {', '.join(sorted(SUPPORTED))}"
        )
    for n, i in enumerate(range(0, len(text), PAGE_CHARS), start=1):
        yield n, text[i : i + PAGE_CHARS]


def ingest_file(path, progress=None) -> int:
    """Chunk, embed and upsert one file. Returns number of chunks stored."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    index = None
    batch, total = [], 0

    def flush():
        nonlocal batch, total, index
        if not batch:
            return
        vectors = embed([b[1] for b in batch])
        if len(vectors) != len(batch):
            raise RuntimeError("Embedding service returned an unexpected vector count")
        if index is None:
            index = ensure_index(len(vectors[0]))
        index.upsert(
            vectors=[
                {"id": b[0], "values": v, "metadata": b[2]}
                for b, v in zip(batch, vectors)
            ]
        )
        total += len(batch)
        batch = []

    for page_number, text in read_pages(path):
        for part, content in enumerate(chunks(text), start=1):
            key = hashlib.sha256(
                f"{path.name}:{page_number}:{part}:{content}".encode()
            ).hexdigest()
            batch.append(
                (
                    key,
                    content,
                    {"page": page_number, "text": content, "file": path.name},
                )
            )
            if len(batch) >= BATCH:
                flush()
        if progress and page_number % 50 == 0:
            progress(page_number)
    flush()
    if total == 0:
        raise ValueError("No readable text found (scanned PDFs need OCR first).")
    return total


def main():
    parser = argparse.ArgumentParser(description="Index files into the knowledge base")
    parser.add_argument("files", nargs="*", help="files to ingest")
    parser.add_argument("--pdf", help="(old option) single PDF path")
    args = parser.parse_args()
    files = args.files or ([args.pdf] if args.pdf else ["data/medical_book.pdf"])
    for f in files:
        print(f"Indexing {f} ...", flush=True)
        n = ingest_file(
            f, progress=lambda p: print(f"  processed {p} pages", flush=True)
        )
        print(f"  done: {n} chunks")
    print("Ingestion complete")


if __name__ == "__main__":
    main()
